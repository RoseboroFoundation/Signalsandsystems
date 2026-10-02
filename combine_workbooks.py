"""Combine all three essay workbooks into signs_and_signals_complete.xlsx."""
import openpyxl
import copy
import os

BASE = os.path.dirname(os.path.abspath(__file__))

sources = [
    ("E1_", os.path.join(BASE, "essay1_full_results.xlsx")),
    ("E2_", os.path.join(BASE, "essay2_results.xlsx")),
    ("E3_", os.path.join(BASE, "essay3_full_results.xlsx")),
]

out_path = os.path.join(BASE, "signs_and_signals_complete.xlsx")
combined = openpyxl.Workbook()
combined.remove(combined.active)  # remove default sheet

for prefix, path in sources:
    print(f"Loading {os.path.basename(path)}...")
    wb = openpyxl.load_workbook(path)
    for sheet_name in wb.sheetnames:
        src = wb[sheet_name]
        new_name = f"{prefix}{sheet_name}"[:31]  # Excel 31-char limit
        dst = combined.create_sheet(title=new_name)

        # Copy cell values and basic formatting
        for row in src.iter_rows():
            for cell in row:
                new_cell = dst.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    new_cell.font = copy.copy(cell.font)
                    new_cell.border = copy.copy(cell.border)
                    new_cell.fill = copy.copy(cell.fill)
                    new_cell.number_format = cell.number_format
                    new_cell.protection = copy.copy(cell.protection)
                    new_cell.alignment = copy.copy(cell.alignment)

        # Copy column widths
        for col_letter, dim in src.column_dimensions.items():
            dst.column_dimensions[col_letter].width = dim.width

        # Copy merged cells
        for merged in src.merged_cells.ranges:
            dst.merge_cells(str(merged))

        # Copy images
        for img in src._images:
            try:
                dst.add_image(copy.copy(img))
            except Exception:
                pass

        print(f"  {new_name}")

    wb.close()

print(f"\nSaving {out_path}...")
combined.save(out_path)
size_mb = os.path.getsize(out_path) / (1024 * 1024)
print(f"Done! {combined.sheetnames.__len__()} sheets, {size_mb:.1f} MB")
