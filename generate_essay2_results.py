"""
Generate Excel workbook with all Essay 2 results (tables and figures).
Reads from the DataStore (SQLite/AWS) and writes to essay2_results.xlsx.
"""
import sys
import os
import warnings
warnings.filterwarnings('ignore')

# Add project to path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList

# ── Helpers ──────────────────────────────────────────────────────────

HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(bold=True, size=12, italic=True)
SIG_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
THIN_BORDER = Border(
    bottom=Side(style='thin', color='CCCCCC'),
)
NUM_FMT_4 = '0.0000'
NUM_FMT_2 = '0.00'
NUM_FMT_6 = '0.000000'
PCT_FMT = '0.0%'


def write_df_to_sheet(ws, df, start_row=1, start_col=1, title=None, subtitle=None):
    """Write a DataFrame to a worksheet with formatted headers."""
    row = start_row

    if title:
        ws.cell(row=row, column=start_col, value=title).font = TITLE_FONT
        row += 1
    if subtitle:
        ws.cell(row=row, column=start_col, value=subtitle).font = SUBTITLE_FONT
        row += 1
    if title or subtitle:
        row += 1  # blank row

    if df.empty:
        ws.cell(row=row, column=start_col, value="No data available")
        return row + 1

    # Headers
    for c_idx, col_name in enumerate(df.columns, start_col):
        cell = ws.cell(row=row, column=c_idx, value=str(col_name))
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    header_row = row
    row += 1

    # Data
    for r_idx, (_, data_row) in enumerate(df.iterrows()):
        for c_idx, (col_name, value) in enumerate(data_row.items(), start_col):
            cell = ws.cell(row=row + r_idx, column=c_idx)
            if pd.isna(value):
                cell.value = ""
            elif isinstance(value, (np.floating, float)):
                cell.value = float(value)
                if 'P_VALUE' in str(col_name) or 'P_' in str(col_name):
                    cell.number_format = NUM_FMT_4
                elif 'PCT' in str(col_name):
                    cell.number_format = NUM_FMT_2
                elif 'COEFF' in str(col_name) or 'CAR' in str(col_name) or 'MEAN' in str(col_name):
                    cell.number_format = NUM_FMT_6
                else:
                    cell.number_format = NUM_FMT_4
            elif isinstance(value, (np.integer, int)):
                cell.value = int(value)
            elif isinstance(value, bool) or isinstance(value, np.bool_):
                cell.value = bool(value)
            else:
                cell.value = str(value)
            cell.border = THIN_BORDER

    end_row = row + len(df)

    # Auto-width columns
    for c_idx, col_name in enumerate(df.columns, start_col):
        max_len = len(str(col_name))
        for r in range(row, end_row):
            val = ws.cell(row=r, column=c_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[ws.cell(row=1, column=c_idx).column_letter].width = min(max_len + 3, 30)

    return end_row + 2


def add_significance_stars(df, p_col='P_VALUE'):
    """Add a SIG column with *, **, *** based on p-value."""
    if p_col not in df.columns:
        return df
    df = df.copy()
    def _sig(p):
        if pd.isna(p): return ''
        if p < 0.01: return '***'
        if p < 0.05: return '**'
        if p < 0.10: return '*'
        return ''
    df['SIG'] = df[p_col].apply(_sig)
    return df


# ── Main ─────────────────────────────────────────────────────────────

def main():
    from model.datastore import DataStore

    print("Loading DataStore...")
    store = DataStore()

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # ================================================================
    # ESSAY 2 NLP (essay2.py) RESULTS
    # ================================================================

    # ── Table 1: News Sentiment Summary ──
    print("  Reading ESSAY2_NEWS_SENTIMENT...")
    news_df = store.read_table('ESSAY2_NEWS_SENTIMENT')
    ws = wb.create_sheet("T1 - News Sentiment")
    if not news_df.empty:
        # Summary by ticker
        summary = news_df.groupby('TICKER').agg(
            N_ARTICLES=('TICKER', 'count'),
            MEAN_SENTIMENT=('SENTIMENT', 'mean'),
            MEAN_WEIGHTED=('SENT_WEIGHTED', 'mean'),
            PCT_POSITIVE=('FINBERT_LABEL', lambda x: (x == 'positive').mean()),
            PCT_NEGATIVE=('FINBERT_LABEL', lambda x: (x == 'negative').mean()),
            PCT_NEUTRAL=('FINBERT_LABEL', lambda x: (x == 'neutral').mean()),
            MEAN_CONFIDENCE=('FINBERT_CONF', 'mean'),
        ).reset_index().round(4)
        write_df_to_sheet(ws, summary,
                          title="Table 1: FinBERT News Sentiment by Company",
                          subtitle="Culture war news articles scored with ProsusAI/FinBERT")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 1: FinBERT News Sentiment by Company")

    # ── Table 2: Filing Sentiment ──
    print("  Reading ESSAY2_FILING_SENTIMENT...")
    filing_df = store.read_table('ESSAY2_FILING_SENTIMENT')
    ws = wb.create_sheet("T2 - Filing Sentiment")
    if not filing_df.empty:
        filing_summary = filing_df.drop(columns=['RUN_TIMESTAMP'], errors='ignore')
        write_df_to_sheet(ws, filing_summary,
                          title="Table 2: SEC Filing Sentiment (10-K/10-Q)",
                          subtitle="MD&A and Risk Factors sections scored with FinBERT")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 2: SEC Filing Sentiment (10-K/10-Q)")

    # ── Table 3: Event NLP Panel ──
    print("  Reading ESSAY2_EVENT_NLP...")
    event_nlp = store.read_table('ESSAY2_EVENT_NLP')
    ws = wb.create_sheet("T3 - Event NLP Panel")
    if not event_nlp.empty:
        cols = [c for c in event_nlp.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, event_nlp[cols],
                          title="Table 3: Event-Window NLP Features",
                          subtitle="Pre/post news sentiment + filing tone per (ticker, event)")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 3: Event-Window NLP Features")

    # ── Table 4: Political Alignment Scores ──
    print("  Reading ESSAY2_POLITICAL_ALIGNMENT...")
    alignment_df = store.read_table('ESSAY2_POLITICAL_ALIGNMENT')
    ws = wb.create_sheet("T4 - Political Alignment")
    if not alignment_df.empty:
        display_cols = [c for c in alignment_df.columns
                        if c not in ('RUN_TIMESTAMP',)]
        write_df_to_sheet(ws, alignment_df[display_cols].round(4),
                          title="Table 4: Political Alignment Scores",
                          subtitle="Three-signal composite: distinctive phrases + stance + cosine")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 4: Political Alignment Scores")

    # ── Table 5: Distinctive Phrases ──
    print("  Reading ESSAY2_DISTINCTIVE_PHRASES...")
    phrases_df = store.read_table('ESSAY2_DISTINCTIVE_PHRASES')
    ws = wb.create_sheet("T5 - Distinctive Phrases")
    if not phrases_df.empty:
        cols = [c for c in phrases_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, phrases_df[cols],
                          title="Table 5: Distinctive Platform Phrases",
                          subtitle="TF-IDF discriminating terms between R and D platforms")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 5: Distinctive Platform Phrases")

    # ── Table 6: Alignment Validation ──
    print("  Reading ESSAY2_ALIGNMENT_VALIDATION...")
    val_df = store.read_table('ESSAY2_ALIGNMENT_VALIDATION')
    ws = wb.create_sheet("T6 - Alignment Validation")
    if not val_df.empty:
        cols = [c for c in val_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, val_df[cols].round(4),
                          title="Table 6: Alignment Validation vs Hand-Coded Leaning",
                          subtitle="Computed vs ESTIMATED_POLITICAL_LEANING comparison")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 6: Alignment Validation vs Hand-Coded Leaning")

    # ── Table 7: Threshold Sensitivity ──
    print("  Reading ESSAY2_THRESHOLD_SENSITIVITY...")
    sens_df = store.read_table('ESSAY2_THRESHOLD_SENSITIVITY')
    ws = wb.create_sheet("T7 - Threshold Sensitivity")
    if not sens_df.empty:
        write_df_to_sheet(ws, sens_df.round(4),
                          title="Table 7: Classification Threshold Sensitivity",
                          subtitle="How alignment classification changes with threshold tau")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 7: Classification Threshold Sensitivity")

    # ── Table 8: NLP Covariate Balance ──
    print("  Reading ESSAY2_NLP_COVARIATE_BALANCE...")
    balance_nlp = store.read_table('ESSAY2_NLP_COVARIATE_BALANCE')
    ws = wb.create_sheet("T8 - NLP Covariate Balance")
    if not balance_nlp.empty:
        write_df_to_sheet(ws, balance_nlp.round(4),
                          title="Table 8: NLP Covariate Balance (Pre-Treatment)",
                          subtitle="Standardized mean differences: treatment vs control")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 8: NLP Covariate Balance (Pre-Treatment)")

    # ── Table 9: Alignment Config ──
    print("  Reading ESSAY2_ALIGNMENT_CONFIG...")
    config_df = store.read_table('ESSAY2_ALIGNMENT_CONFIG')
    ws = wb.create_sheet("T9 - Alignment Config")
    if not config_df.empty:
        write_df_to_sheet(ws, config_df,
                          title="Table 9: Alignment Run Configuration",
                          subtitle="Thresholds, weights, and agreement rate")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 9: Alignment Run Configuration")

    # ================================================================
    # ESSAY 2 DiD (essay2_did.py) RESULTS
    # ================================================================

    # ── Table 10: CAR Panel ──
    print("  Reading ESSAY2_CAR_PANEL...")
    car_panel = store.read_table('ESSAY2_CAR_PANEL')
    ws = wb.create_sheet("T10 - CAR Panel")
    if not car_panel.empty:
        cols = [c for c in car_panel.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, car_panel[cols].round(6),
                          title="Table 10: Cumulative Abnormal Returns Panel",
                          subtitle="(Firm, Event) CARs from FF5 normal-return model")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 10: Cumulative Abnormal Returns Panel")

    # ── Table 11: DiD Coefficient Table ──
    print("  Reading ESSAY2_DID_COEFFICIENTS...")
    coeff_df = store.read_table('ESSAY2_DID_COEFFICIENTS')
    ws = wb.create_sheet("T11 - DiD Coefficients")
    if not coeff_df.empty:
        cols = [c for c in coeff_df.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(coeff_df[cols])
        write_df_to_sheet(ws, display.round(6),
                          title="Table 11: DiD Regression Coefficients",
                          subtitle="Three specs: Basic, With Lean, With FOMO")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 11: DiD Regression Coefficients")

    # ── Table 12: Parallel Trends ──
    print("  Reading ESSAY2_PARALLEL_TRENDS...")
    pt_df = store.read_table('ESSAY2_PARALLEL_TRENDS')
    ws = wb.create_sheet("T12 - Parallel Trends")
    if not pt_df.empty:
        cols = [c for c in pt_df.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(pt_df[cols])
        row_end = write_df_to_sheet(ws, display.round(6),
                          title="Table 12: Parallel Trends Pre-Test (Treatment vs Control)",
                          subtitle="Day-by-day Treat x Day coefficients in [-10, -1]")

        # Figure: Parallel Trends Plot
        if len(pt_df) >= 3 and 'DAY' in pt_df.columns and 'COEFFICIENT' in pt_df.columns:
            chart = LineChart()
            chart.title = "Figure 1: Pre-Event Parallel Trends"
            chart.x_axis.title = "Trading Day Relative to Event"
            chart.y_axis.title = "Treat x Day Coefficient"
            chart.style = 10
            chart.width = 20
            chart.height = 12

            # Write chart data
            chart_start = row_end + 2
            ws.cell(row=chart_start, column=1, value="DAY").font = HEADER_FONT
            ws.cell(row=chart_start, column=2, value="COEFFICIENT").font = HEADER_FONT
            for i, (_, r) in enumerate(pt_df.sort_values('DAY').iterrows()):
                ws.cell(row=chart_start + 1 + i, column=1, value=int(r['DAY']))
                ws.cell(row=chart_start + 1 + i, column=2, value=float(r['COEFFICIENT']))

            data = Reference(ws, min_col=2, min_row=chart_start, max_row=chart_start + len(pt_df))
            cats = Reference(ws, min_col=1, min_row=chart_start + 1, max_row=chart_start + len(pt_df))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, f"A{chart_start + len(pt_df) + 3}")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 12: Parallel Trends Pre-Test")

    # ── Table 13: Multi-Window Summary ──
    print("  Reading ESSAY2_MULTI_WINDOW_SUMMARY...")
    mw_summary = store.read_table('ESSAY2_MULTI_WINDOW_SUMMARY')
    ws = wb.create_sheet("T13 - Multi-Window Summary")
    if not mw_summary.empty:
        cols = [c for c in mw_summary.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(mw_summary[cols], 'P_VALUE_VS_ZERO')
        row_end = write_df_to_sheet(ws, display.round(4),
                          title="Table 13: Multi-Window Event Study Summary",
                          subtitle="Treatment CARs by window [-1, +N]")

        # Figure: Multi-Window CAR Bar Chart
        if len(mw_summary) >= 2 and 'MEAN_CAR_TREAT' in mw_summary.columns:
            chart = BarChart()
            chart.title = "Figure 2: Treatment Mean CARs by Post-Event Window"
            chart.x_axis.title = "Event Window"
            chart.y_axis.title = "Mean CAR"
            chart.style = 10
            chart.width = 20
            chart.height = 12

            chart_start = row_end + 2
            ws.cell(row=chart_start, column=1, value="WINDOW").font = HEADER_FONT
            ws.cell(row=chart_start, column=2, value="MEAN_CAR_TREAT").font = HEADER_FONT
            ws.cell(row=chart_start, column=3, value="MEAN_CAR_CTRL").font = HEADER_FONT
            for i, (_, r) in enumerate(mw_summary.iterrows()):
                ws.cell(row=chart_start + 1 + i, column=1, value=str(r.get('WINDOW', '')))
                ws.cell(row=chart_start + 1 + i, column=2, value=float(r['MEAN_CAR_TREAT']) if not pd.isna(r['MEAN_CAR_TREAT']) else 0)
                ctrl_val = r.get('MEAN_CAR_CTRL', np.nan)
                ws.cell(row=chart_start + 1 + i, column=3, value=float(ctrl_val) if not pd.isna(ctrl_val) else 0)

            data = Reference(ws, min_col=2, max_col=3, min_row=chart_start, max_row=chart_start + len(mw_summary))
            cats = Reference(ws, min_col=1, min_row=chart_start + 1, max_row=chart_start + len(mw_summary))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, f"A{chart_start + len(mw_summary) + 3}")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 13: Multi-Window Event Study Summary")

    # ── Table 14: Treatment vs Control ──
    print("  Reading ESSAY2_MULTI_WINDOW_TREAT_VS_CTRL...")
    tvc_df = store.read_table('ESSAY2_MULTI_WINDOW_TREAT_VS_CTRL')
    ws = wb.create_sheet("T14 - Treat vs Control")
    if not tvc_df.empty:
        cols = [c for c in tvc_df.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(tvc_df[cols])
        write_df_to_sheet(ws, display.round(4),
                          title="Table 14: Treatment vs Control (Welch t-test)",
                          subtitle="Per-window treatment-control CAR differences")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 14: Treatment vs Control (Welch t-test)")

    # ── Table 15: CARs by Political Lean ──
    print("  Reading ESSAY2_MULTI_WINDOW_BY_LEAN...")
    lean_df = store.read_table('ESSAY2_MULTI_WINDOW_BY_LEAN')
    ws = wb.create_sheet("T15 - CARs by Lean")
    if not lean_df.empty:
        cols = [c for c in lean_df.columns if c != 'RUN_TIMESTAMP']
        row_end = write_df_to_sheet(ws, lean_df[cols].round(4),
                          title="Table 15: Treatment CARs by Political Lean",
                          subtitle="Conservative vs Liberal vs Mixed across windows")

        # Figure: CARs by Lean
        if not lean_df.empty and 'LEAN' in lean_df.columns and 'MEAN_CAR' in lean_df.columns:
            pivot = lean_df.pivot_table(index='WINDOW', columns='LEAN', values='MEAN_CAR')
            if not pivot.empty:
                chart = BarChart()
                chart.type = "col"
                chart.grouping = "clustered"
                chart.title = "Figure 3: Treatment CARs by Political Lean"
                chart.x_axis.title = "Window"
                chart.y_axis.title = "Mean CAR"
                chart.style = 10
                chart.width = 22
                chart.height = 12

                chart_start = row_end + 2
                lean_cols = [c for c in ['Conservative', 'Liberal', 'Mixed'] if c in pivot.columns]
                ws.cell(row=chart_start, column=1, value="WINDOW").font = HEADER_FONT
                for ci, lc in enumerate(lean_cols, 2):
                    ws.cell(row=chart_start, column=ci, value=lc).font = HEADER_FONT

                for i, (win, vals) in enumerate(pivot.iterrows()):
                    ws.cell(row=chart_start + 1 + i, column=1, value=str(win))
                    for ci, lc in enumerate(lean_cols, 2):
                        v = vals.get(lc, 0)
                        ws.cell(row=chart_start + 1 + i, column=ci, value=float(v) if not pd.isna(v) else 0)

                data = Reference(ws, min_col=2, max_col=1 + len(lean_cols), min_row=chart_start, max_row=chart_start + len(pivot))
                cats = Reference(ws, min_col=1, min_row=chart_start + 1, max_row=chart_start + len(pivot))
                chart.add_data(data, titles_from_data=True)
                chart.set_categories(cats)
                ws.add_chart(chart, f"A{chart_start + len(pivot) + 3}")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 15: Treatment CARs by Political Lean")

    # ── Table 15a: Headline Descriptive Stats (Item 1) ──
    print("  Reading ESSAY2_DESCRIPTIVE_STATS...")
    desc_df = store.read_table('ESSAY2_DESCRIPTIVE_STATS')
    ws = wb.create_sheet("T15a - Descriptive Stats")
    if not desc_df.empty:
        cols = [c for c in desc_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, desc_df[cols].round(6),
                          title="Table 15a: Headline Descriptive Statistics",
                          subtitle="EW/VW Mean CAR, BHAR, Median, Significance — all event windows [-1, +5] through [-1, +90]")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 15a: Headline Descriptive Statistics")

    # ── Table 15b: Per-Lean Significance Tests (Item 2) ──
    print("  Reading ESSAY2_MULTI_WINDOW_LEAN_TESTS...")
    lean_tests = store.read_table('ESSAY2_MULTI_WINDOW_LEAN_TESTS')
    ws = wb.create_sheet("T15b - Lean Significance")
    if not lean_tests.empty:
        cols = [c for c in lean_tests.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(lean_tests[cols])
        write_df_to_sheet(ws, display.round(4),
                          title="Table 15b: Per-Lean One-Sample t-Tests (CARs != 0)",
                          subtitle="Conservative, Liberal, Mixed — each tested separately")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 15b: Per-Lean Significance Tests")

    # ── Table 15c: Pairwise Lean Comparisons (Item 2) ──
    print("  Reading ESSAY2_MULTI_WINDOW_LEAN_PAIRWISE...")
    lean_pw = store.read_table('ESSAY2_MULTI_WINDOW_LEAN_PAIRWISE')
    ws = wb.create_sheet("T15c - Lean Pairwise")
    if not lean_pw.empty:
        cols = [c for c in lean_pw.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(lean_pw[cols])
        write_df_to_sheet(ws, display.round(4),
                          title="Table 15c: Pairwise Lean Comparisons (Welch t-test)",
                          subtitle="Conservative-Liberal gap, Mixed-Liberal, Mixed-Conservative")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 15c: Pairwise Lean Comparisons")

    # ── Table 15d: Cross-Sectional Regression (Item 3) ──
    print("  Reading ESSAY2_CROSS_SECTIONAL_REG...")
    xsect_df = store.read_table('ESSAY2_CROSS_SECTIONAL_REG')
    ws = wb.create_sheet("T15d - XSect Regression")
    if not xsect_df.empty:
        cols = [c for c in xsect_df.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(xsect_df[cols])
        write_df_to_sheet(ws, display.round(6),
                          title="Table 15d: Cross-Sectional Regression (Safe-Haven / Ambiguity)",
                          subtitle="CAR_POST ~ ALIGNMENT_SCORE + VIX_LEVEL + FOMO_Z")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 15d: Cross-Sectional Regression")

    # ── Table 16: Contagion Summary ──
    print("  Reading ESSAY2_CONTAGION_SUMMARY...")
    cont_summary = store.read_table('ESSAY2_CONTAGION_SUMMARY')
    ws = wb.create_sheet("T16 - Contagion Summary")
    if not cont_summary.empty:
        cols = [c for c in cont_summary.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(cont_summary[cols], 'P_VALUE_VS_ZERO')
        write_df_to_sheet(ws, display.round(4),
                          title="Table 16: Industry Contagion — Peer CARs vs Zero",
                          subtitle="Do industry peers experience abnormal returns?")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 16: Industry Contagion — Peer CARs vs Zero")

    # ── Table 17: Peer vs Non-Peer ──
    print("  Reading ESSAY2_CONTAGION_PEER_VS_NONPEER...")
    pvnp_df = store.read_table('ESSAY2_CONTAGION_PEER_VS_NONPEER')
    ws = wb.create_sheet("T17 - Peer vs Non-Peer")
    if not pvnp_df.empty:
        cols = [c for c in pvnp_df.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(pvnp_df[cols])
        write_df_to_sheet(ws, display.round(4),
                          title="Table 17: Peer vs Non-Peer Differential Contagion",
                          subtitle="Are peer CARs different from non-peer CARs?")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 17: Peer vs Non-Peer Differential Contagion")

    # ── Table 18: Contagion by Lean ──
    print("  Reading ESSAY2_CONTAGION_BY_LEAN...")
    cont_lean = store.read_table('ESSAY2_CONTAGION_BY_LEAN')
    ws = wb.create_sheet("T18 - Contagion by Lean")
    if not cont_lean.empty:
        cols = [c for c in cont_lean.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, cont_lean[cols].round(4),
                          title="Table 18: Peer Contagion by Triggering Firm's Lean",
                          subtitle="Conservative vs Liberal vs Mixed event triggers")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 18: Peer Contagion by Triggering Firm's Lean")

    # ── Table 19: Tight Non-Peer Differential ──
    print("  Reading ESSAY2_CONTAGION_TIGHT_DIFF...")
    tight_df = store.read_table('ESSAY2_CONTAGION_TIGHT_DIFF')
    ws = wb.create_sheet("T19 - Tight Non-Peer Diff")
    if not tight_df.empty:
        cols = [c for c in tight_df.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(tight_df[cols])
        write_df_to_sheet(ws, display.round(4),
                          title="Table 19: Tight Non-Peer Differential (Different NAICS Sector)",
                          subtitle="Enhanced contagion: peers vs firms in different 2-digit sectors")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 19: Tight Non-Peer Differential")

    # ── Table 20: Consumer vs B2B Contagion ──
    print("  Reading ESSAY2_CONTAGION_BY_FACING...")
    facing_df = store.read_table('ESSAY2_CONTAGION_BY_FACING')
    ws = wb.create_sheet("T20 - Consumer vs B2B")
    if not facing_df.empty:
        cols = [c for c in facing_df.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(facing_df[cols], 'P_VALUE_VS_ZERO')
        write_df_to_sheet(ws, display.round(4),
                          title="Table 20: Contagion by Industry Type (Consumer vs B2B)",
                          subtitle="Does contagion hit harder in consumer-facing industries?")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 20: Contagion by Industry Type")

    # ── Table 21: Consumer vs B2B Tests ──
    print("  Reading ESSAY2_CONTAGION_CONS_VS_B2B...")
    cvb_df = store.read_table('ESSAY2_CONTAGION_CONS_VS_B2B')
    ws = wb.create_sheet("T21 - Cons vs B2B Tests")
    if not cvb_df.empty:
        cols = [c for c in cvb_df.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(cvb_df[cols])
        write_df_to_sheet(ws, display.round(4),
                          title="Table 21: Consumer vs B2B Pairwise t-Tests",
                          subtitle="Direct comparison of contagion magnitude")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 21: Consumer vs B2B Pairwise t-Tests")

    # ── Table 22: Lean Mechanism ──
    print("  Reading ESSAY2_CONTAGION_LEAN_MECH...")
    lean_mech = store.read_table('ESSAY2_CONTAGION_LEAN_MECH')
    ws = wb.create_sheet("T22 - Lean Mechanism")
    if not lean_mech.empty:
        cols = [c for c in lean_mech.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(lean_mech[cols], 'P_VALUE_VS_ZERO')
        write_df_to_sheet(ws, display.round(4),
                          title="Table 22: Contagion Lean Mechanism (Uncertainty Channel)",
                          subtitle="Peer CARs by triggering firm's lean with significance")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 22: Contagion Lean Mechanism")

    # ── Table 23: Lean Pairwise Tests ──
    print("  Reading ESSAY2_CONTAGION_LEAN_PAIRWISE...")
    lean_pw = store.read_table('ESSAY2_CONTAGION_LEAN_PAIRWISE')
    ws = wb.create_sheet("T23 - Lean Pairwise Tests")
    if not lean_pw.empty:
        cols = [c for c in lean_pw.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(lean_pw[cols])
        write_df_to_sheet(ws, display.round(4),
                          title="Table 23: Lean Pairwise Comparisons",
                          subtitle="Mixed vs Liberal, Mixed vs Conservative, Liberal vs Conservative")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 23: Lean Pairwise Comparisons")

    # ── Table 24: Peer Parallel Trends ──
    print("  Reading ESSAY2_PEER_PARALLEL_TRENDS...")
    peer_pt = store.read_table('ESSAY2_PEER_PARALLEL_TRENDS')
    ws = wb.create_sheet("T24 - Peer Parallel Trends")
    if not peer_pt.empty:
        cols = [c for c in peer_pt.columns if c != 'RUN_TIMESTAMP']
        display = add_significance_stars(peer_pt[cols])
        write_df_to_sheet(ws, display.round(6),
                          title="Table 24: Peer Parallel Trends (Contagion Validity)",
                          subtitle="Event firm vs industry peers, pre-event [-10, -1]")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 24: Peer Parallel Trends")

    # ================================================================
    # DIAGNOSTIC TESTS
    # ================================================================

    # ── Table 25: Placebo Test ──
    print("  Reading ESSAY2_DIAG_PLACEBO...")
    placebo = store.read_table('ESSAY2_DIAG_PLACEBO')
    ws = wb.create_sheet("T25 - Placebo Test")
    if not placebo.empty:
        # Show summary + distribution stats
        actual = placebo[placebo.get('ROW_TYPE', pd.Series()) == 'ACTUAL_SUMMARY']
        placebo_only = placebo[placebo.get('ROW_TYPE', pd.Series()) == 'PLACEBO']

        summary_data = []
        if not actual.empty:
            a = actual.iloc[0]
            summary_data.append({
                'METRIC': 'Actual TREAT_x_POST Coefficient',
                'VALUE': a.get('PLACEBO_COEFF', np.nan),
            })
            summary_data.append({
                'METRIC': 'Percentile Rank (fraction of placebos more extreme)',
                'VALUE': a.get('PLACEBO_P', np.nan),
            })
        if not placebo_only.empty:
            summary_data.append({
                'METRIC': 'N Placebo Iterations',
                'VALUE': len(placebo_only),
            })
            summary_data.append({
                'METRIC': 'Placebo Mean Coefficient',
                'VALUE': placebo_only['PLACEBO_COEFF'].mean(),
            })
            summary_data.append({
                'METRIC': 'Placebo Std Coefficient',
                'VALUE': placebo_only['PLACEBO_COEFF'].std(),
            })

        row_end = write_df_to_sheet(ws, pd.DataFrame(summary_data).round(6),
                          title="Table 25: Placebo Test (Randomization Inference)",
                          subtitle="Permute treatment assignment, re-estimate DiD")

        # Figure: Placebo distribution
        if not placebo_only.empty and 'PLACEBO_COEFF' in placebo_only.columns:
            # Histogram data (bins)
            coeffs = placebo_only['PLACEBO_COEFF'].dropna().values
            if len(coeffs) > 10:
                hist, edges = np.histogram(coeffs, bins=30)
                chart_start = row_end + 2
                ws.cell(row=chart_start, column=1, value="BIN_CENTER").font = HEADER_FONT
                ws.cell(row=chart_start, column=2, value="COUNT").font = HEADER_FONT
                for i in range(len(hist)):
                    center = (edges[i] + edges[i+1]) / 2
                    ws.cell(row=chart_start + 1 + i, column=1, value=round(center, 6))
                    ws.cell(row=chart_start + 1 + i, column=2, value=int(hist[i]))

                chart = BarChart()
                chart.title = "Figure 4: Placebo Coefficient Distribution"
                chart.x_axis.title = "Placebo TREAT_x_POST Coefficient"
                chart.y_axis.title = "Count"
                chart.style = 10
                chart.width = 20
                chart.height = 12

                data = Reference(ws, min_col=2, min_row=chart_start, max_row=chart_start + len(hist))
                cats = Reference(ws, min_col=1, min_row=chart_start + 1, max_row=chart_start + len(hist))
                chart.add_data(data, titles_from_data=True)
                chart.set_categories(cats)
                ws.add_chart(chart, f"A{chart_start + len(hist) + 3}")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 25: Placebo Test")

    # ── Table 26: Bootstrap CI ──
    print("  Reading ESSAY2_DIAG_BOOTSTRAP_CI...")
    boot_df = store.read_table('ESSAY2_DIAG_BOOTSTRAP_CI')
    ws = wb.create_sheet("T26 - Bootstrap CI")
    if not boot_df.empty:
        write_df_to_sheet(ws, boot_df.round(6),
                          title="Table 26: Bootstrap Confidence Intervals",
                          subtitle="Block bootstrap (firm x event) for treatment vs control CAR diff")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 26: Bootstrap Confidence Intervals")

    # ── Table 27: Cluster-Robust SEs ──
    print("  Reading ESSAY2_DIAG_CLUSTER_ROBUST...")
    cluster_df = store.read_table('ESSAY2_DIAG_CLUSTER_ROBUST')
    ws = wb.create_sheet("T27 - Cluster Robust SEs")
    if not cluster_df.empty:
        cols = [c for c in cluster_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, cluster_df[cols].round(6),
                          title="Table 27: Cluster-Robust Standard Errors Comparison",
                          subtitle="Cluster by TICKER vs Cluster by EVENT_ID")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 27: Cluster-Robust SEs Comparison")

    # ── Table 28: Normality Tests ──
    print("  Reading ESSAY2_DIAG_NORMALITY...")
    norm_df = store.read_table('ESSAY2_DIAG_NORMALITY')
    ws = wb.create_sheet("T28 - Normality Tests")
    if not norm_df.empty:
        cols = [c for c in norm_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, norm_df[cols].round(4),
                          title="Table 28: Residual Normality Tests",
                          subtitle="Jarque-Bera and Shapiro-Wilk on DiD residuals")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 28: Residual Normality Tests")

    # ── Table 29: Heteroskedasticity ──
    print("  Reading ESSAY2_DIAG_HETEROSKEDASTICITY...")
    het_df = store.read_table('ESSAY2_DIAG_HETEROSKEDASTICITY')
    ws = wb.create_sheet("T29 - Heteroskedasticity")
    if not het_df.empty:
        cols = [c for c in het_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, het_df[cols].round(4),
                          title="Table 29: Heteroskedasticity Tests",
                          subtitle="Breusch-Pagan and White tests")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 29: Heteroskedasticity Tests")

    # ── Table 30: VIF ──
    print("  Reading ESSAY2_DIAG_VIF...")
    vif_df = store.read_table('ESSAY2_DIAG_VIF')
    ws = wb.create_sheet("T30 - VIF")
    if not vif_df.empty:
        cols = [c for c in vif_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, vif_df[cols].round(2),
                          title="Table 30: Variance Inflation Factors",
                          subtitle="VIF > 10 indicates multicollinearity")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 30: Variance Inflation Factors")

    # ── Table 31: Covariate Balance ──
    print("  Reading ESSAY2_DIAG_COVARIATE_BALANCE...")
    bal_df = store.read_table('ESSAY2_DIAG_COVARIATE_BALANCE')
    ws = wb.create_sheet("T31 - Covariate Balance")
    if not bal_df.empty:
        cols = [c for c in bal_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, bal_df[cols].round(4),
                          title="Table 31: DiD Covariate Balance (Treatment vs Control)",
                          subtitle="Standardized Mean Differences; |SMD| < 0.10 = balanced")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 31: DiD Covariate Balance")

    # ── Table 32: Autocorrelation ──
    print("  Reading ESSAY2_DIAG_AUTOCORRELATION...")
    ac_df = store.read_table('ESSAY2_DIAG_AUTOCORRELATION')
    ws = wb.create_sheet("T32 - Autocorrelation")
    if not ac_df.empty:
        cols = [c for c in ac_df.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, ac_df[cols].round(3),
                          title="Table 32: Durbin-Watson Autocorrelation Test",
                          subtitle="DW ~ 2 = no autocorrelation")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 32: Durbin-Watson Autocorrelation Test")

    # ── Table 33: Event Alignment Scores ──
    print("  Reading ESSAY2_EVENT_ALIGNMENT...")
    evt_align = store.read_table('ESSAY2_EVENT_ALIGNMENT')
    ws = wb.create_sheet("T33 - Event Alignment")
    if not evt_align.empty:
        cols = [c for c in evt_align.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, evt_align[cols].round(4),
                          title="Table 33: Event-Level Political Alignment Scores",
                          subtitle="Alignment score per (ticker, event_date)")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 33: Event-Level Political Alignment Scores")

    # ── Table 34: Multi-Window Full Panel ──
    print("  Reading ESSAY2_MULTI_WINDOW_PANEL...")
    mw_panel = store.read_table('ESSAY2_MULTI_WINDOW_PANEL')
    ws = wb.create_sheet("T34 - MW Full Panel")
    if not mw_panel.empty:
        cols = [c for c in mw_panel.columns if c != 'RUN_TIMESTAMP']
        write_df_to_sheet(ws, mw_panel[cols].round(6),
                          title="Table 34: Multi-Window Full Panel (Long Format)",
                          subtitle="(Ticker, Event, Window) CARs for all firms")
    else:
        write_df_to_sheet(ws, pd.DataFrame(),
                          title="Table 34: Multi-Window Full Panel")

    # ── Save ──
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'essay2_results.xlsx')
    wb.save(output_path)
    store.close()
    print(f"\nSaved: {output_path}")
    print(f"Sheets: {len(wb.sheetnames)}")
    for name in wb.sheetnames:
        print(f"  - {name}")


if __name__ == '__main__':
    main()
